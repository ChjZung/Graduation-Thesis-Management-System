import sys
import json
import os
import base64
import io
import docx

def format_score(val):
    if val is None:
        return ""
    try:
        f = float(val)
        return f"{f:.2f}"
    except (ValueError, TypeError):
        return ""

def main():
    if len(sys.argv) < 4:
        print("Usage: python export_phieu_cham_word.py <template_path> <json_path> <output_path>")
        sys.exit(1)

    template_path = sys.argv[1]
    json_path = sys.argv[2]
    output_path = sys.argv[3]

    if not os.path.exists(template_path):
        print(f"Error: Template not found at {template_path}")
        sys.exit(1)

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    doc = docx.Document(template_path)

    ten_de_tai = data.get('ten_de_tai', '')
    sinh_viens = data.get('sinh_viens', [])
    gvhd_ho_ten = data.get('gvhd_ho_ten', '')
    ten_hoi_dong = data.get('ten_hoi_dong', '')
    phong = data.get('phong', '')
    ngay_cham = data.get('ngay_cham', '')
    custom_chuc_nang = data.get('custom_chuc_nang', '')
    scores = data.get('scores', {})  # map: ma_sv -> {"tc1": 0.25, ..., "total": 9.0}
    giang_vien_ho_ten = data.get('giang_vien_ho_ten', '')
    vai_tro_gv = data.get('vai_tro_gv', 'Giám khảo chấm điểm')
    loai_khoa_luan = data.get('loai_khoa_luan', 'KLCN')

    is_klks = (loai_khoa_luan == 'KLKS')

    # 1. Cập nhật các đoạn văn bản thông tin đầu trang
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith('Tên đề tài:'):
            p.text = f"Tên đề tài: {ten_de_tai}"
        elif txt.startswith('1. Mã SV:'):
            if len(sinh_viens) >= 1:
                sv = sinh_viens[0]
                p.text = f"1. Mã SV: {sv.get('ma_sv', '')}\tHọ tên: {sv.get('ho_ten', '')}\tLớp: {sv.get('lop', '')}"
            else:
                p.text = "1. Mã SV:\tHọ tên:\tLớp:"
        elif txt.startswith('2. Mã SV:'):
            if len(sinh_viens) >= 2:
                sv = sinh_viens[1]
                p.text = f"2. Mã SV: {sv.get('ma_sv', '')}\tHọ tên: {sv.get('ho_ten', '')}\tLớp: {sv.get('lop', '')}"
            else:
                p.text = ""
        elif txt.startswith('3. Mã SV:'):
            if len(sinh_viens) >= 3:
                sv = sinh_viens[2]
                p.text = f"3. Mã SV: {sv.get('ma_sv', '')}\tHọ tên: {sv.get('ho_ten', '')}\tLớp: {sv.get('lop', '')}"
            else:
                p.text = ""
        elif txt.startswith('Giảng viên hướng dẫn:'):
            p.text = f"Giảng viên hướng dẫn: {gvhd_ho_ten}"
        elif txt.startswith('Hội đồng số:') or txt.startswith('Hội đồng:'):
            p.text = f"Hội đồng: {ten_hoi_dong}\nPhòng: {phong}\tNgày chấm: {ngay_cham}"

    # 2. Cập nhật bảng điểm
    if len(doc.tables) > 0:
        table = doc.tables[0]

        # Cập nhật tên sinh viên ở hàng sub-header (Row 1)
        sub_header_row = table.rows[1]
        for s_idx, sv in enumerate(sinh_viens[:3]):
            col_idx = 4 + s_idx
            if col_idx < len(sub_header_row.cells):
                cell = sub_header_row.cells[col_idx]
                cell.text = f"SV{s_idx + 1}\n{sv.get('ho_ten', '')}"

        # Ánh xạ hàng tới tiêu chí
        # Row 2 -> tc1, Row 3 -> tc2, Row 4 -> tc3, Row 5 -> tc4, Row 6 -> tc5, Row 7 -> tc6, Row 8 -> tc7
        criteria_keys_klcn = ['tc1', 'tc2', 'tc3', 'tc4', 'tc5', 'tc6', 'tc7', 'tc8', 'tc9', 'tc10', 'tc11', 'tc12', 'tc13']
        criteria_keys_klks = ['tc1', 'tc2', 'tc3', 'tc4', 'tc5', 'tc6', 'tc7', 'tc8', 'tc9', 'tc10', 'tc11', 'tc12', 'tc13', 'tc14']

        crit_keys = criteria_keys_klks if is_klks else criteria_keys_klcn

        # Cập nhật nội dung mục chức năng bài toán (Row 8 - tc7)
        row_tc7_idx = 8
        if row_tc7_idx < len(table.rows):
            cell_tc7_content = table.rows[row_tc7_idx].cells[1]
            if custom_chuc_nang and custom_chuc_nang.strip():
                if is_klks:
                    cell_tc7_content.text = f"Xây dựng ứng dụng:\n- Xây dựng các chức năng cơ bản của ứng dụng: Đăng nhập, đăng xuất, đổi mật khẩu, phân quyền ứng dụng\n- Cài đặt chức năng nghiệp vụ theo yêu cầu bài toán:\n{custom_chuc_nang.strip()}"
                else:
                    cell_tc7_content.text = f"Xây dựng các chức năng cơ bản của ứng dụng: đăng nhập, đăng xuất, backup & restore,…\nCài đặt chức năng theo yêu cầu của bài toán\n{custom_chuc_nang.strip()}"

        # Điền điểm cho từng tiêu chí
        for idx, c_key in enumerate(crit_keys):
            r_idx = 2 + idx
            if r_idx >= len(table.rows) - 1:
                break
            row = table.rows[r_idx]
            for s_idx, sv in enumerate(sinh_viens[:3]):
                col_idx = 4 + s_idx
                if col_idx < len(row.cells):
                    sv_scores = scores.get(sv.get('ma_sv', ''), {})
                    sc = sv_scores.get(c_key)
                    row.cells[col_idx].text = format_score(sc)

        # Hàng tổng cộng (hàng cuối cùng)
        last_row = table.rows[-1]
        for s_idx, sv in enumerate(sinh_viens[:3]):
            col_idx = 4 + s_idx
            if col_idx < len(last_row.cells):
                sv_scores = scores.get(sv.get('ma_sv', ''), {})
                tot = sv_scores.get('total')
                last_row.cells[col_idx].text = format_score(tot)

    # 3. Ký tên và họ tên giảng viên ở cuối trang
    # Thêm thông tin ngày tháng và tên giám khảo
    p_sign = doc.add_paragraph()
    p_sign.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
    run_date = p_sign.add_run(f"\nTP. Hồ Chí Minh, ngày {ngay_cham}\n")
    run_date.italic = True
    run_title = p_sign.add_run("GIÁM KHẢO CHẤM ĐIỂM\n")
    run_title.bold = True
    run_note = p_sign.add_run("(Ký và ghi rõ họ tên)\n")
    run_note.italic = True

    # Chèn ảnh chữ ký vẽ điện tử nếu có
    chu_ky_b64 = payload.get('chu_ky')
    if chu_ky_b64 and 'base64,' in chu_ky_b64:
        try:
            b64_data = chu_ky_b64.split('base64,')[1]
            img_bytes = base64.b64decode(b64_data)
            img_stream = io.BytesIO(img_bytes)
            p_sign.add_run().add_picture(img_stream, width=docx.shared.Inches(1.8))
            p_sign.add_run("\n")
        except Exception as e:
            p_sign.add_run("\n\n\n\n")
    else:
        p_sign.add_run("\n\n\n\n")

    run_name = p_sign.add_run(f"{giang_vien_ho_ten}\n")
    run_name.bold = True
    run_role = p_sign.add_run(f"({vai_tro_gv})")
    run_role.font.size = docx.shared.Pt(9.5)

    doc.save(output_path)
    print(f"SUCCESS: Saved word document to {output_path}")

if __name__ == '__main__':
    main()
